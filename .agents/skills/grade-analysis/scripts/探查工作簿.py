# -*- coding: utf-8 -*-
"""
探查成绩表 xlsx 的结构，为写 配置.json 提供依据。

用法：
    python 探查工作簿.py <xlsx路径> [工作表名]

不给工作表名则列出全部工作表。给出工作表名则进一步推断：
表头行、列分组（每三列一组）、学号样式与班别位次、等级集合、异常值。
"""
import re
import sys
from collections import Counter
from openpyxl import load_workbook


def cell_str(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def col_letter(i):
    """1 -> A, 27 -> AA"""
    s = ""
    while i > 0:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def list_sheets(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    print(f"工作簿：{path}")
    print(f"{'工作表名':<16}{'状态':<8}{'行数':>7}{'列数':>7}")
    for ws in wb.worksheets:
        print(f"{ws.title:<16}{ws.sheet_state:<8}{ws.max_row:>7}{ws.max_column:>7}")
    print("\n提示：隐藏的工作表（state=hidden）往往是拆分结果，不是数据源；"
          "数据源通常是含学号与各科成绩的那张。")
    wb.close()


def probe_sheet(path, sheet):
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise SystemExit(f"没有工作表「{sheet}」。现有：{wb.sheetnames}")
    ws = wb[sheet]
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        rows.append(list(row))
        if i >= 30:
            break
    wb.close()

    print(f"工作表：{sheet}")
    print(f"前 30 行中最大列数：{max((len(r) for r in rows), default=0)}")
    print()

    # 1) 找表头行：前 10 行里非空单元格最多的那一行
    def filled(r):
        return sum(1 for v in r if cell_str(v) != "")
    head_idx = max(range(min(10, len(rows))), key=lambda i: filled(rows[i]))
    header = rows[head_idx]
    print(f"== 表头行推断：第 {head_idx + 1} 行（非空 {filled(header)} 格）==")
    preview = [f"{col_letter(i+1)}:{cell_str(v)}" for i, v in enumerate(header) if cell_str(v)]
    print("  " + "  ".join(preview[:30]))
    print()

    # 2) 列分组：自第 3 列起每三列一组（分数/排名/等第）
    print("== 列分组推断（自第 3 列起，每 3 列一组）==")
    groups, i = [], 2
    while i + 2 < len(header):
        name = cell_str(header[i]) or f"(第{i+1}列)"
        groups.append((name, i, i + 1, i + 2))
        i += 3
    for name, a, b, c in groups:
        print(f"  {name:<8} 分数={col_letter(a+1):<4} 排名={col_letter(b+1):<4} 等第={col_letter(c+1):<4}"
              f"  表头文字：{cell_str(header[a])}/{cell_str(header[b])}/{cell_str(header[c])}")
    if (len(header) - 2) % 3:
        print(f"  ⚠ 第 3 列起共 {len(header)-2} 列，不是 3 的整数倍 —— 分组假设可能不成立，需人工确认")
    print()

    # 3) 前两列
    print("== 前两列前 8 行 ==")
    for r in rows[head_idx + 1:head_idx + 9]:
        print("  " + " | ".join(cell_str(v) for v in r[:2]))
    print()

    # 4) 学号与班别位次
    ids = []
    for r in rows[head_idx + 1:]:
        s = cell_str(r[0]) if r else ""
        if s:
            ids.append(s)
    lens = Counter(len(x) for x in ids)
    print(f"== 学号（A 列）长度分布：{dict(lens)}  样本：{ids[:3]}")
    digit_ids = [x for x in ids if re.fullmatch(r"\d+", x)]
    if digit_ids and len(lens) == 1:
        L = len(digit_ids[0])
        print(f"   全部为 {L} 位数字学号，按不同起始位切班别编号：")
        for start in range(1, L):
            width = min(2, L - start + 1)
            seg = Counter(x[start - 1:start - 1 + width] for x in digit_ids)
            kind = "← 像班别" if 6 <= len(seg) <= 20 else ""
            print(f"     第 {start}~{start+width-1} 位 -> {len(seg):>3} 个取值 "
                  f"{sorted(seg)[:12]} {kind}")
    print()

    # 5) 各等第列实际出现的值
    print("== 各科目「等第」列出现的值 ==")
    for name, a, b, c in groups:
        vals = Counter()
        blanks = 0
        for r in rows[head_idx + 1:]:
            v = cell_str(r[c]) if c < len(r) else ""
            if v:
                vals[v] += 1
            else:
                blanks += 1
        keys = sorted(vals)
        print(f"  {name:<8} {len(keys):>2} 种：{''.join(keys)}"
              + (f"   （另有 {blanks} 条为空）" if blanks else ""))
    print()

    print("== 提醒 ==")
    print("  · 若等第列出现字母表外的值（如 S/T 之后还有），脚本会报错提示，按实际补进「等级顺序」")
    print("  · 若学号长度不统一或含非数字，拆分脚本会归入 未分班.csv，需人工核对")
    print("  · 「排名」列是否唯一、是否从 1 开始，决定了能否用「排名基准科目」推导等级区间")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    if len(sys.argv) >= 3:
        probe_sheet(sys.argv[1], sys.argv[2])
    else:
        list_sheets(sys.argv[1])
