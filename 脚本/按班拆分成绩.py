# -*- coding: utf-8 -*-
"""
按班别拆分成绩表 → 每个班一份 CSV。

用法：
    python 按班拆分成绩.py <考试目录>

    <考试目录> 例如 九年级9月考，也可以给完整路径。
    目录内需有成绩表 xlsx；结果写到该目录下的 分班数据/。

目录结构参数来自 <考试目录>/配置.json，文件不存在时用下面的默认值。
**新考试的表格格式若与默认不同，请先用「成绩表分析」skill 探查，再写配置.json。**
"""
import csv
import json
import os
import re
import shutil
import sys
import tempfile
from collections import OrderedDict

from openpyxl import load_workbook

# ==================== 默认配置 ====================
# 与「表结构」相关；行号列号均从 1 开始数，和 Excel 里看到的一致
DEFAULT_CONFIG = {
    "源文件": None,          # None = 目录里唯一的 xlsx；多个时须在此写明文件名
    "工作表": "成绩分析",     # 含学号与各科成绩的表
    "表头行": 1,             # 表头所在行
    "数据起始行": 2,          # 第一条学生记录所在行
    "学号列": 1,             # 学号所在列（A=1）
    "班别位起": 6,           # 学号中班别的起始位次（032420102 的 01 在第 6~7 位）
    "班别位数": 2,
    "保留列数": 26,          # 只保留前 N 列（A~Z），之后的列忽略
}

# 成绩表里的班级编号集合；用于核对拆分结果是否覆盖完整
EXPECTED_CLASSES = ["01", "02", "03", "04", "05", "06",
                    "07", "08", "09", "10", "11", "12"]

# 本脚本不读、但同属考试配置的键（分析脚本要用）；不应报成“未使用”
ANALYSIS_KEYS = {"重点班级", "等级顺序", "排名基准科目"}


def resolve_exam_dir(arg):
    """把命令行参数解析成考试目录的绝对路径。"""
    here = os.path.dirname(os.path.abspath(__file__))
    ws = os.path.dirname(here)                      # 成绩分析工作区/
    cand = arg if os.path.isabs(arg) else os.path.join(ws, arg)
    if not os.path.isdir(cand):
        cand2 = os.path.join(os.getcwd(), arg)
        if os.path.isdir(cand2):
            cand = cand2
        else:
            raise SystemExit(f"找不到考试目录：{arg}\n  已尝试：{cand}")
    return os.path.abspath(cand)


def read_config(path):
    """读 配置.json。用 utf-8-sig：记事本 / PowerShell 保存的 UTF-8 会带 BOM，
    普通 utf-8 读取会直接抛 JSONDecodeError。"""
    try:
        with open(path, encoding="utf-8-sig") as f:
            return json.load(f)
    except json.JSONDecodeError as e:
        raise SystemExit(f"配置.json 不是合法的 JSON：{path}\n  {e}\n"
                         f"  常见原因：多余的逗号、缺引号、中文引号「」而非英文 \"\"。")
    except UnicodeDecodeError:
        raise SystemExit(f"配置.json 编码不是 UTF-8：{path}\n"
                         f"  请用记事本另存为 UTF-8 编码后重试。")


def load_config(exam_dir):
    cfg = dict(DEFAULT_CONFIG)
    path = os.path.join(exam_dir, "配置.json")
    if os.path.exists(path):
        user = read_config(path)
        # 下划线开头的键（如 "_说明"）视为注释，不参与校验
        unknown = {k for k in user if not k.startswith("_")} - set(DEFAULT_CONFIG) - ANALYSIS_KEYS
        if unknown:
            print(f"⚠ 配置.json 里有无法识别的键（拼写错误？）：{sorted(unknown)}")
        cfg.update(user)
        print(f"已读取配置：{path}")
    else:
        print("未找到 配置.json，使用默认结构（成绩分析表 / 表头第 1 行 / 学号在 A 列）")
    return cfg


def pick_source(exam_dir, cfg):
    """确定要读的 xlsx。"""
    if cfg["源文件"]:
        p = os.path.join(exam_dir, cfg["源文件"])
        if not os.path.exists(p):
            raise SystemExit(f"配置指定的源文件不存在：{p}")
        return p
    xlsx = [f for f in os.listdir(exam_dir)
            if f.lower().endswith((".xlsx", ".xlsm")) and not f.startswith("~$")]
    if not xlsx:
        raise SystemExit(f"目录内没有 xlsx：{exam_dir}")
    if len(xlsx) > 1:
        raise SystemExit("目录内有多个 xlsx，请在 配置.json 里用「源文件」指定要用哪个：\n  "
                         + "\n  ".join(sorted(xlsx)))
    return os.path.join(exam_dir, xlsx[0])


def normalize_id(value):
    """学号规整为纯数字字符串；非数字返回 None。"""
    if value is None:
        return None
    if isinstance(value, float):
        if not value.is_integer():
            return None
        value = int(value)
    text = str(value).strip()
    return text if re.fullmatch(r"\d+", text) else None


def is_blank_row(values):
    """整行无实际内容（全 None 或空串），用于跳过残留的空公式行。"""
    return all(v is None or (isinstance(v, str) and v.strip() == "") for v in values)


def clean_value(value):
    """单元格值转成适合写 CSV 的字符串。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def main():
    if len(sys.argv) < 2:
        raise SystemExit(__doc__.strip().split("用法：")[1].split("\n")[0].strip()
                         + "\n\n请给出考试目录名，例如：python 按班拆分成绩.py 九年级9月考")

    exam_dir = resolve_exam_dir(sys.argv[1])
    cfg = load_config(exam_dir)
    src = pick_source(exam_dir, cfg)
    out_dir = os.path.join(exam_dir, "分班数据")

    print(f"考试目录：{exam_dir}")
    print(f"源文件　：{os.path.basename(src)}")
    print(f"工作表　：{cfg['工作表']}")

    # 文件可能正被 Excel 打开占用，先复制临时副本
    tmp_path = None
    try:
        wb = load_workbook(src, read_only=True, data_only=True)
    except PermissionError:
        fd, tmp_path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        shutil.copy2(src, tmp_path)
        print("提示：源文件正被占用，已复制临时副本读取。")
        wb = load_workbook(tmp_path, read_only=True, data_only=True)

    try:
        if cfg["工作表"] not in wb.sheetnames:
            raise SystemExit(
                f"工作簿中没有名为「{cfg['工作表']}」的工作表。\n"
                f"现有工作表：{', '.join(wb.sheetnames)}\n"
                f"若名称不同，请在 配置.json 里改「工作表」。")
        ws = wb[cfg["工作表"]]

        keep = cfg["保留列数"]
        header = None
        header_full = None      # 未截断的表头，用于检查「保留列数」是否切掉了内容
        groups = OrderedDict()
        row_no = 0
        for row in ws.iter_rows(values_only=True):
            row_no += 1
            if row_no < cfg["数据起始行"] and row_no != cfg["表头行"]:
                continue
            if row is None or is_blank_row(row):
                continue
            values = list(row)[:keep]
            values += [None] * (keep - len(values))
            if row_no == cfg["表头行"]:
                header = values
                header_full = list(row)     # 完整表头（未截断），供越界检查用
                continue
            sid = normalize_id(values[cfg["学号列"] - 1])
            if sid is None:
                cls = "未分班"
            else:
                sid = sid.zfill(9)     # 数字型学号补回前导零：32421062 -> 032421062
                b, n = cfg["班别位起"] - 1, cfg["班别位数"]
                cls = sid[b:b + n] if len(sid) >= b + n else "未分班"
            groups.setdefault(cls, []).append(values)
    finally:
        wb.close()
        if tmp_path and os.path.exists(tmp_path):
            os.remove(tmp_path)

    if header is None:
        raise SystemExit(f"第 {cfg['表头行']} 行没有读到表头，请检查 配置.json 里的行号。")

    # 各年级的表列数不一样，且可能远超 A~Z（已见 26/30/32/35/36 五种）。
    # 防一手「保留列数」把真实列切掉：在被截断的范围之外，若还存在完整的科目组
    # （某列有表头文字、且其后两列恰为「排名」「等第」），说明保留列数给小了。
    # 注意不能用「表头非空单元格数」判断——有的表在数据列之后还摆了图例，
    # 其表头行也带文字，会被误判成数据列（九年级表的 AA~AC 就是这种）。
    tail = header_full
    beyond = [i for i in range(keep, max(0, len(tail) - 2))
              if str(tail[i] or "").strip()
              and str(tail[i + 1] or "").strip() == "排名"
              and str(tail[i + 2] or "").strip() == "等第"]
    if beyond:
        raise SystemExit(
            f"⚠「保留列数」= {keep} 太小：第 {beyond[0] + 1} 列起还有完整的科目组"
            f"（表头「{str(tail[beyond[0]]).strip()}」+ 排名 + 等第），会被丢弃。\n"
            f"  请把 配置.json 的「保留列数」改为 {beyond[-1] + 3} 后重跑。")

    os.makedirs(out_dir, exist_ok=True)
    total = 0
    for cls, rows in groups.items():
        name = "未分班.csv" if cls == "未分班" else f"{cls}班.csv"
        with open(os.path.join(out_dir, name), "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow([clean_value(v) for v in header])
            for r in rows:
                w.writerow([clean_value(v) for v in r])
        total += len(rows)

    # 自检
    print()
    print("=== 拆分结果 ===")
    for cls in sorted(groups, key=lambda c: (c == "未分班", c)):
        print(f"  {cls:>4s}班 {len(groups[cls]):>4d} 人"
              if cls != "未分班" else f"  未分班 {len(groups[cls]):>4d} 人")
    print(f"  合计 {total} 人，输出目录：{out_dir}")

    missing = [c for c in EXPECTED_CLASSES if c not in groups]
    if missing:
        print(f"  提示：以下班级本次没有数据：{', '.join(missing)}")
    extra = [c for c in groups if c not in EXPECTED_CLASSES and c != "未分班"]
    if extra:
        print(f"  ⚠ 出现预期外的班别编号：{', '.join(sorted(extra))}"
              f" —— 请核对 配置.json 的「班别位起/班别位数」是否正确。")
    if "未分班" in groups:
        print(f"  ⚠ 有 {len(groups['未分班'])} 人学号为空或格式异常，见 未分班.csv")


if __name__ == "__main__":
    main()
